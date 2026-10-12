"""Bounded file inspector executed by the shared execution runtime."""

from __future__ import annotations

import csv
import json
import sys
from itertools import islice
from pathlib import Path
from typing import cast

_PREVIEW_CHARS = 12_000
_MAX_JSON_PARSE_BYTES = 2 * 1024 * 1024


def inspect_file(path: Path) -> dict[str, object]:
    suffix = path.suffix.lower()
    result: dict[str, object] = {"filename": path.name, "size_bytes": path.stat().st_size}
    if suffix in {".txt", ".md"}:
        result["preview"] = path.read_text(encoding="utf-8")[:_PREVIEW_CHARS]
    elif suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8", newline="") as source:
            rows = list(islice(csv.reader(source, delimiter="\t" if suffix == ".tsv" else ","), 21))
        bounded = [[_bounded_cell(cell) for cell in row[:20]] for row in rows]
        result.update({"columns": bounded[0] if bounded else [], "rows": bounded[1:21]})
    elif suffix == ".json":
        result["preview"] = path.read_text(encoding="utf-8")[:_PREVIEW_CHARS]
        result["structure"] = (
            _structure(json.loads(path.read_text(encoding="utf-8")))
            if path.stat().st_size <= _MAX_JSON_PARSE_BYTES
            else {"type": "large_json", "parsed": False}
        )
    elif suffix == ".parquet":
        import pyarrow.parquet as pq  # type: ignore[import-untyped]

        parquet = pq.ParquetFile(path)
        first_batch = next(parquet.iter_batches(batch_size=10), None)
        parquet_rows = cast(list[dict[str, object]], first_batch.to_pylist() if first_batch else [])
        result.update({
            "schema": str(parquet.schema_arrow)[:8_000],
            "row_count": parquet.metadata.num_rows,
            "rows": [
                {str(key): _bounded_cell(value) for key, value in list(row.items())[:20]}
                for row in parquet_rows
            ],
        })
    elif suffix == ".xlsx":
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        workbook = load_workbook(path, read_only=True, data_only=True)
        result["sheets"] = [
            {
                "name": sheet.title,
                "rows": [
                    [_bounded_cell(cell) for cell in row[:15]]
                    for row in sheet.iter_rows(max_row=10, max_col=15, values_only=True)
                ],
            }
            for sheet in workbook.worksheets[:3]
        ]
        workbook.close()
    elif suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        result["paragraphs"] = [_bounded_cell(item.text, 500) for item in document.paragraphs[:80]]
        result["table_count"] = len(document.tables)
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(path))
        result["slides"] = [
            [_bounded_cell(shape.text, 300) for shape in slide.shapes if hasattr(shape, "text")][
                :10
            ]
            for slide in presentation.slides[:10]
        ]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        result["page_count"] = len(reader.pages)
        result["preview"] = "\n".join(
            (page.extract_text() or "")[:4_000] for page in reader.pages[:5]
        )[:_PREVIEW_CHARS]
    elif suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        with Image.open(path) as image:
            result.update({
                "width": image.width,
                "height": image.height,
                "mode": image.mode,
                "format": image.format,
            })
    return result


def _bounded_cell(value: object, limit: int = 120) -> object:
    if value is None or isinstance(value, int | float | bool):
        return value
    text = str(value)
    return text if len(text) <= limit else f"{text[:limit]}…"


def _structure(value: object, depth: int = 0) -> object:
    if depth >= 3:
        return type(value).__name__
    if isinstance(value, dict):
        return {str(key): _structure(item, depth + 1) for key, item in list(value.items())[:50]}
    if isinstance(value, list):
        return {
            "type": "array",
            "length": len(value),
            "sample": [_structure(item, depth + 1) for item in value[:3]],
        }
    return type(value).__name__


if __name__ == "__main__":
    print(json.dumps(inspect_file(Path(sys.argv[1])), ensure_ascii=False, default=str))

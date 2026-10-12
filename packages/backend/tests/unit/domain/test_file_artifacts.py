from __future__ import annotations

import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from rhetorica_backend.contracts.chat import AnalystChatAskRequest, ChatActorContext, ChatAskRequest
from rhetorica_backend.contracts.execution import QueryPreview
from rhetorica_backend.contracts.files import FileArtifactRef
from rhetorica_backend.domain.chat.llm.prompt_builder import PromptBuilder
from rhetorica_backend.domain.chat.orchestration.state import ToolOrchestrationState
from rhetorica_backend.domain.chat.orchestration.tool_file_capability_service import (
    ToolFileCapabilityService,
)
from rhetorica_backend.domain.chat.orchestration.tool_result_prompt_serializer import (
    ToolResultPromptSerializer,
)
from rhetorica_backend.domain.chat.orchestration.transport_runtime_services import (
    StreamingResponseContentService,
)
from rhetorica_backend.domain.execution.snapshots import SqlResultSnapshot
from rhetorica_backend.domain.files.models import (
    Artifact,
    ArtifactSource,
    ArtifactStatus,
    SandboxInput,
    SandboxOutput,
    SandboxRunResult,
    StoredObject,
)
from rhetorica_backend.domain.files.policy import (
    ArtifactPolicy,
    ArtifactPolicyConfig,
    ArtifactRejectedError,
)
from rhetorica_backend.domain.files.ports import SandboxGatewayError
from rhetorica_backend.domain.files.service import ArtifactService
from rhetorica_backend.domain.files.sql_export import (
    SqlResultExportAccessError,
    SqlResultExportLimitError,
    SqlResultExportService,
)
from rhetorica_backend.infra.execution import ProcessExecutor, ProcessLimits
from rhetorica_backend.infra.execution.file_sessions import (
    ExecutionRejectedError,
    FileSessionConfig,
    FileSessionRuntime,
)
from rhetorica_backend.infra.object_storage import ObjectStorageSqlResultSnapshotStore
from rhetorica_backend.infra.object_storage.sql_result_snapshots import (
    SqlResultSnapshotLimitError,
)
from rhetorica_backend.infra.postgres.chat_repository import ChatRepository
from rhetorica_backend.infra.postgres.execution_history_repository import (
    ExecutionHistoryRepository,
)
from rhetorica_backend.infra.postgres.file_artifact_repository import FileArtifactRepository
from rhetorica_backend.infra.postgres.file_models import FileArtifactOrm, FileQuotaUsageOrm
from rhetorica_backend.infra.postgres.session import DatabaseRuntime
from rhetorica_backend.infra.sandbox.http_gateway import HttpSandboxGateway
from sqlalchemy import event
from sqlalchemy.orm import Session


def _ooxml(marker: str, payload: bytes = b"ok") -> io.BytesIO:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr(marker, payload)
    output.seek(0)
    return output


def test_policy_validates_magic_paths_and_ooxml_expansion() -> None:
    policy = ArtifactPolicy()
    assert policy.normalize_filename("report.pdf") == "report.pdf"
    assert (
        policy.inspect(
            filename="report.pdf",
            declared_media_type="application/pdf",
            fileobj=io.BytesIO(b"%PDF-1.7\n"),
        )
        == "application/pdf"
    )

    with pytest.raises(ArtifactRejectedError, match="Invalid file name"):
        policy.normalize_filename("../report.pdf")
    with pytest.raises(ArtifactRejectedError) as mismatch:
        policy.inspect(
            filename="report.pdf",
            declared_media_type="application/pdf",
            fileobj=io.BytesIO(b"not a pdf"),
        )
    assert mismatch.value.code == "magic_mismatch"

    strict_policy = ArtifactPolicy(ArtifactPolicyConfig(max_ooxml_compression_ratio=2.0))
    with pytest.raises(ArtifactRejectedError) as bomb:
        strict_policy.inspect(
            filename="bomb.xlsx",
            declared_media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            fileobj=_ooxml("xl/workbook.xml", b"0" * 100_000),
        )
    assert bomb.value.code == "archive_ratio"
    trusted_generated = _ooxml("xl/workbook.xml", b"0" * 100_000)
    assert (
        strict_policy.inspect(
            filename="generated.xlsx",
            declared_media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            fileobj=trusted_generated,
            trusted_generated=True,
        )
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    macro_file = _ooxml("xl/workbook.xml", b"ok")
    with zipfile.ZipFile(macro_file, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"macro")
    macro_file.seek(0)
    with pytest.raises(ArtifactRejectedError) as macros:
        policy.inspect(
            filename="macro.xlsx",
            declared_media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            fileobj=macro_file,
        )
    assert macros.value.code == "macros_forbidden"


def test_policy_reads_all_text_before_accepting_utf8() -> None:
    policy = ArtifactPolicy()
    invalid_after_first_chunk = io.BytesIO(b"a" * (1024 * 1024 + 1) + b"\xff")

    with pytest.raises(ArtifactRejectedError) as invalid:
        policy.inspect(
            filename="late-invalid.csv",
            declared_media_type="text/csv",
            fileobj=invalid_after_first_chunk,
        )

    assert invalid.value.code == "invalid_text_encoding"


def test_chat_request_requires_text_or_attachment() -> None:
    with pytest.raises(ValueError):
        AnalystChatAskRequest(question="")
    request = AnalystChatAskRequest(question="", attachment_ids=["artifact-id"])
    assert request.attachment_ids == ["artifact-id"]


def test_repository_enforces_ownership_lineage_and_message_lifecycle(tmp_path: Path) -> None:
    runtime = DatabaseRuntime(f"sqlite:///{tmp_path}/files.db")
    runtime.create_schema()
    chat = ChatRepository(runtime)
    conversation = chat.create_conversation(workspace_id="default", title="Files")
    repository = FileArtifactRepository(runtime)
    artifact = repository.create(
        artifact_id="artifact-1",
        workspace_id="default",
        owner_user_id=conversation.owner_user_id,
        conversation_id=conversation.id,
        storage_key="objects/artifact-1",
        filename="input.csv",
        media_type="text/csv",
        size_bytes=8,
        sha256="a" * 64,
        source=ArtifactSource.USER_UPLOAD,
        status=ArtifactStatus.STAGED,
    )
    assert repository.get_owned(artifact.id, "somebody-else", "default") is None
    assert repository.get_owned(artifact.id, conversation.owner_user_id, "other") is None

    message = chat.save_message(
        conversation_id=conversation.id,
        role="user",
        content="Analyze",
        owner_user_id=conversation.owner_user_id,
    )
    linked = repository.attach_to_message(
        artifact_ids=[artifact.id],
        message_id=message.id,
        conversation_id=conversation.id,
        owner_user_id=conversation.owner_user_id,
        workspace_id="default",
        relation="input",
    )
    assert linked[0].status is ArtifactStatus.READY
    derived = repository.create(
        artifact_id="artifact-2",
        workspace_id="default",
        owner_user_id=conversation.owner_user_id,
        conversation_id=conversation.id,
        storage_key="objects/artifact-2",
        filename="output.csv",
        media_type="text/csv",
        size_bytes=9,
        sha256="b" * 64,
        source=ArtifactSource.SANDBOX_OUTPUT,
        status=ArtifactStatus.STAGED,
        source_artifact_ids=[artifact.id],
    )
    assert derived.source_artifact_ids == (artifact.id,)
    with pytest.raises(ValueError, match="lineage ownership"):
        repository.create(
            artifact_id="artifact-foreign-lineage",
            workspace_id="default",
            owner_user_id=conversation.owner_user_id,
            conversation_id=conversation.id,
            storage_key="objects/artifact-foreign-lineage",
            filename="output.csv",
            media_type="text/csv",
            size_bytes=9,
            sha256="c" * 64,
            source=ArtifactSource.SANDBOX_OUTPUT,
            status=ArtifactStatus.STAGED,
            source_artifact_ids=["not-owned"],
        )
    assistant_message = chat.save_message(
        conversation_id=conversation.id,
        role="assistant",
        content="Done",
        owner_user_id=conversation.owner_user_id,
    )
    repository.attach_to_message(
        artifact_ids=[derived.id],
        message_id=assistant_message.id,
        conversation_id=conversation.id,
        owner_user_id=conversation.owner_user_id,
        workspace_id="default",
        relation="output",
    )
    assert (
        repository.delete_staged(
            artifact_id=artifact.id,
            owner_user_id=conversation.owner_user_id,
            workspace_id="default",
        )
        is None
    )
    history = chat.get_conversation_with_messages(
        conversation.id, owner_user_id=conversation.owner_user_id
    )
    assert history is not None
    assert history.messages[0].files[0].id == artifact.id
    assert history.messages[1].files[0].id == derived.id
    assert repository.usage_bytes(workspace_id="default") == 17

    assert chat.delete_conversation(conversation.id, owner_user_id=conversation.owner_user_id)
    assert (
        repository.get_owned(artifact.id, conversation.owner_user_id, "default").status
        is ArtifactStatus.READY
    )
    assert repository.usage_bytes(workspace_id="default") == 17


def test_attach_and_delete_lock_artifact_rows_for_atomic_status_transition(tmp_path: Path) -> None:
    runtime = DatabaseRuntime(f"sqlite:///{tmp_path}/artifact-transitions.db")
    runtime.create_schema()
    chat = ChatRepository(runtime)
    conversation = chat.create_conversation(workspace_id="default", title="Transitions")
    message = chat.save_message(
        conversation_id=conversation.id,
        role="user",
        content="Attach",
        owner_user_id=conversation.owner_user_id,
    )
    repository = FileArtifactRepository(runtime)
    for artifact_id in ("artifact-attach", "artifact-delete"):
        repository.create(
            artifact_id=artifact_id,
            workspace_id="default",
            owner_user_id=conversation.owner_user_id,
            conversation_id=conversation.id,
            storage_key=f"objects/{artifact_id}",
            filename=f"{artifact_id}.csv",
            media_type="text/csv",
            size_bytes=1,
            sha256="a" * 64,
            source=ArtifactSource.USER_UPLOAD,
            status=ArtifactStatus.STAGED,
        )

    locked_file_queries = 0

    def record_lock(execute_state) -> None:
        nonlocal locked_file_queries
        statement = execute_state.statement
        if (
            execute_state.is_select
            and getattr(statement, "_for_update_arg", None) is not None
            and any(
                description.get("entity") is FileArtifactOrm
                for description in getattr(statement, "column_descriptions", [])
            )
        ):
            locked_file_queries += 1

    event.listen(Session, "do_orm_execute", record_lock)
    try:
        repository.attach_to_message(
            artifact_ids=["artifact-attach"],
            message_id=message.id,
            conversation_id=conversation.id,
            owner_user_id=conversation.owner_user_id,
            workspace_id="default",
            relation="input",
        )
        repository.delete_staged(
            artifact_id="artifact-delete",
            owner_user_id=conversation.owner_user_id,
            workspace_id="default",
        )
    finally:
        event.remove(Session, "do_orm_execute", record_lock)

    assert locked_file_queries == 2


def test_conversation_soft_delete_retains_sql_snapshot(tmp_path: Path) -> None:
    runtime = DatabaseRuntime(f"sqlite:///{tmp_path}/snapshot-cleanup.db")
    runtime.create_schema()
    chat = ChatRepository(runtime)
    conversation = chat.create_conversation(workspace_id="default", title="Snapshot")
    history = ExecutionHistoryRepository(runtime)
    history.create_running(
        execution_id="execution-cleanup",
        workspace_id="default",
        engine="clickhouse",
        submitted_sql="SELECT 1",
        normalized_sql="SELECT 1 LIMIT 1",
        timeout_ms=60_000,
        row_limit=1,
        warnings=[],
        violations=[],
        connection_summary={},
        actor_id=conversation.owner_user_id,
    )
    history.complete(
        execution_id="execution-cleanup",
        status="succeeded",
        query_id="query-1",
        row_count=1,
        duration_ms=1,
        error_message=None,
        snapshot=SqlResultSnapshot(
            execution_id="execution-cleanup",
            workspace_id="default",
            owner_user_id=conversation.owner_user_id,
            conversation_id=conversation.id,
            storage_key="snapshots/execution-cleanup.jsonl.gz",
            columns=("value",),
            row_count=1,
            size_bytes=10,
            sha256="a" * 64,
            possibly_truncated=True,
            created_at=datetime.now(tz=UTC),
        ),
    )

    assert chat.delete_conversation(
        conversation.id,
        owner_user_id=conversation.owner_user_id,
    )
    pending = history.list_snapshots_for_cleanup(
        orphan_before=datetime.now(tz=UTC) - timedelta(days=1),
    )
    assert pending == []
    assert history.get_snapshot("execution-cleanup") is not None


class _CaptureArtifactService:
    def __init__(self) -> None:
        self.content = b""

    def store_generated(self, **kwargs):
        self.content = kwargs["fileobj"].read()
        return Artifact(
            id="export-1",
            workspace_id=kwargs["workspace_id"],
            owner_user_id=kwargs["owner_user_id"],
            conversation_id=kwargs["conversation_id"],
            storage_key="hidden",
            filename=kwargs["filename"],
            media_type=kwargs["declared_media_type"],
            size_bytes=len(self.content),
            sha256="b" * 64,
            source=ArtifactSource.SQL_RESULT,
            status=ArtifactStatus.STAGED,
        )

    to_ref = staticmethod(ArtifactService.to_ref)


class _MemoryObjectStorage:
    def __init__(self) -> None:
        self.objects: dict[str, tuple[bytes, str]] = {}

    def upload_fileobj(self, *, key, fileobj, size_bytes, media_type) -> None:
        payload = fileobj.read()
        assert len(payload) == size_bytes
        self.objects[key] = (payload, media_type)

    def open(self, key: str) -> StoredObject:
        payload, media_type = self.objects[key]
        return StoredObject(
            body=io.BytesIO(payload),
            size_bytes=len(payload),
            media_type=media_type,
        )

    def delete(self, key: str) -> None:
        self.objects.pop(key, None)

    def create_download_url(self, key: str, *, expires_seconds: int = 300) -> str:
        return f"memory://{key}?expires={expires_seconds}"


def test_sql_snapshot_limits_uncompressed_bytes_even_when_rows_compress_well() -> None:
    storage = _MemoryObjectStorage()
    snapshots = ObjectStorageSqlResultSnapshotStore(
        storage=storage,  # type: ignore[arg-type]
        max_rows=10,
        max_bytes=1024,
        max_cell_bytes=512,
        max_row_bytes=768,
    )
    capture = snapshots.begin(
        execution_id="execution-compressible",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-1",
    )
    capture.start(["payload"])

    with pytest.raises(SqlResultSnapshotLimitError, match="sql_result_snapshot_size_limit"):
        capture.write_rows([["a" * 400], ["a" * 400], ["a" * 400]])

    capture.abort()
    assert storage.objects == {}


def test_sql_snapshot_rejects_oversized_cell_before_json_serialization(monkeypatch) -> None:
    storage = _MemoryObjectStorage()
    snapshots = ObjectStorageSqlResultSnapshotStore(
        storage=storage,  # type: ignore[arg-type]
        max_rows=1,
        max_bytes=4096,
        max_cell_bytes=1024,
        max_row_bytes=2048,
    )
    capture = snapshots.begin(
        execution_id="execution-large-cell",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-1",
    )
    capture.start(["payload"])

    def unexpected_json_dump(*_args, **_kwargs):
        raise AssertionError("json.dumps must not receive an oversized cell")

    monkeypatch.setattr(
        "rhetorica_backend.infra.object_storage.sql_result_snapshots.json.dumps",
        unexpected_json_dump,
    )
    with pytest.raises(SqlResultSnapshotLimitError, match="sql_result_snapshot_cell_size_limit"):
        capture.write_rows([["a" * 2048]])
    capture.abort()


def test_parallel_uploads_cannot_overcommit_atomic_quota(tmp_path: Path) -> None:
    runtime = DatabaseRuntime(f"sqlite:///{tmp_path}/quota-race.db")
    runtime.create_schema()
    conversation = ChatRepository(runtime).create_conversation(
        workspace_id="default",
        title="Quota race",
    )
    repository = FileArtifactRepository(runtime)
    storage = _MemoryObjectStorage()
    service = ArtifactService(
        repository=repository,
        storage=storage,
        policy=ArtifactPolicy(
            ArtifactPolicyConfig(
                max_input_bytes=100,
                user_quota_bytes=10,
                workspace_quota_bytes=100,
            )
        ),
    )

    def upload(index: int) -> Artifact:
        return service.stage_upload(
            workspace_id="default",
            owner_user_id=conversation.owner_user_id,
            conversation_id=conversation.id,
            filename=f"input-{index}.txt",
            declared_media_type="text/plain",
            fileobj=io.BytesIO(b"123456"),
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(upload, index) for index in range(2)]
    successes = [future.result() for future in futures if future.exception() is None]
    failures = [future.exception() for future in futures if future.exception() is not None]

    assert len(successes) == 1
    assert len(failures) == 1
    assert isinstance(failures[0], ArtifactRejectedError)
    assert failures[0].code == "user_quota_exceeded"
    assert repository.usage_bytes(workspace_id="default") == 6
    assert len(storage.objects) == 1
    with runtime.session() as session:
        usage = {
            (row.scope_type, row.scope_id): row.used_bytes
            for row in session.query(FileQuotaUsageOrm).all()
        }
    assert usage[("workspace", "default")] == 6
    assert usage[("user", conversation.owner_user_id)] == 6


def test_sql_export_uses_persisted_snapshot_and_marks_limit() -> None:
    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=2,
        preview=QueryPreview(columns=["id", "name"], rows=[[1, "one"], [2, "two"]]),
    )
    capture = _CaptureArtifactService()
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(get=lambda execution_id: history),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=capture,  # type: ignore[arg-type]
    )
    result = service.export(
        execution_id="execution-1",
        export_format="csv",
        workspace_id="default",
        conversation_id="conversation-1",
        owner_user_id="user-1",
        allowed_current_execution_ids={"execution-1"},
    )
    assert capture.content == b"id,name\r\n1,one\r\n2,two\r\n"
    assert result.possibly_truncated is True
    assert result.file.kind == "sql_result"


def test_sql_export_allows_only_current_turn_for_legacy_null_actor() -> None:
    history = SimpleNamespace(
        workspace_id="default",
        actor_id=None,
        status="succeeded",
        row_limit=1,
        preview=QueryPreview(columns=["id"], rows=[[1]]),
    )
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(get=lambda execution_id: history),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=_CaptureArtifactService(),  # type: ignore[arg-type]
    )

    with pytest.raises(SqlResultExportAccessError):
        service.export(
            execution_id="legacy-execution",
            export_format="csv",
            workspace_id="default",
            conversation_id="conversation-1",
            owner_user_id="user-1",
        )

    result = service.export(
        execution_id="legacy-execution",
        export_format="csv",
        workspace_id="default",
        conversation_id="conversation-1",
        owner_user_id="user-1",
        allowed_current_execution_ids={"legacy-execution"},
    )
    assert result.row_count == 1


def test_sql_export_rejects_snapshot_from_another_conversation() -> None:
    snapshot = SqlResultSnapshot(
        execution_id="execution-foreign",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-foreign",
        storage_key="snapshot",
        columns=("id",),
        row_count=1,
        size_bytes=10,
        sha256="a" * 64,
        possibly_truncated=False,
        created_at=datetime.now(tz=UTC),
    )
    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=1,
        preview=QueryPreview(columns=["id"], rows=[[1]]),
    )
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(
            get=lambda execution_id: history,
            get_snapshot=lambda execution_id: snapshot,
        ),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=_CaptureArtifactService(),  # type: ignore[arg-type]
        snapshot_store=SimpleNamespace(iter_rows=lambda item: iter([[1]])),
    )

    with pytest.raises(SqlResultExportAccessError):
        service.export(
            execution_id="execution-foreign",
            export_format="xlsx",
            workspace_id="default",
            conversation_id="conversation-1",
            owner_user_id="user-1",
            allowed_current_execution_ids={"execution-foreign"},
        )


def test_sql_export_writes_xlsx_and_neutralizes_formulas() -> None:
    from openpyxl import load_workbook

    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=10,
        preview=QueryPreview(columns=["name"], rows=[['=HYPERLINK("bad")']]),
    )
    capture = _CaptureArtifactService()
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(get=lambda execution_id: history),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=capture,  # type: ignore[arg-type]
    )

    result = service.export(
        execution_id="execution-2",
        export_format="xlsx",
        workspace_id="default",
        conversation_id="conversation-1",
        owner_user_id="user-1",
        allowed_current_execution_ids={"execution-2"},
    )

    workbook = load_workbook(io.BytesIO(capture.content), read_only=True)
    values = list(workbook.active.values)
    assert values == [("name",), ('\'=HYPERLINK("bad")',)]
    assert result.possibly_truncated is False


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        "  =1+1",
        "\t=1+1",
        "\r=1+1",
        "\n=1+1",
        " \t @SUM(A1:A2)",
        "\tplain-control-prefixed-value",
    ],
)
def test_sql_export_neutralizes_formula_injection_after_leading_whitespace(value: str) -> None:
    assert SqlResultExportService._safe_cell(value) == f"'{value}"


def test_sql_export_streams_full_snapshot_instead_of_prompt_preview() -> None:
    storage = _MemoryObjectStorage()
    snapshots = ObjectStorageSqlResultSnapshotStore(
        storage=storage,  # type: ignore[arg-type]
        max_rows=10,
        max_bytes=1024 * 1024,
    )
    capture = snapshots.begin(
        execution_id="execution-full",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-1",
    )
    capture.start(["id", "name"])
    capture.write_rows([[1, "one"], [2, "two"], [3, "three"]])
    snapshot = capture.finish(possibly_truncated=False)
    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=10,
        row_count=3,
        preview=QueryPreview(columns=["id", "name"], rows=[[1, "one"]]),
    )
    artifacts = _CaptureArtifactService()
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(
            get=lambda execution_id: history,
            get_snapshot=lambda execution_id: snapshot,
        ),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=artifacts,  # type: ignore[arg-type]
        snapshot_store=snapshots,
        max_export_rows=10,
    )

    result = service.export(
        execution_id="execution-full",
        export_format="csv",
        workspace_id="default",
        conversation_id="conversation-1",
        owner_user_id="user-1",
        allowed_current_execution_ids={"execution-full"},
    )

    assert result.row_count == 3
    assert artifacts.content == b"id,name\r\n1,one\r\n2,two\r\n3,three\r\n"


def test_sql_export_rejects_manifest_above_configured_row_limit() -> None:
    snapshot = SqlResultSnapshot(
        execution_id="execution-large",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-1",
        storage_key="snapshot",
        columns=("id",),
        row_count=3,
        size_bytes=10,
        sha256="a" * 64,
        possibly_truncated=False,
        created_at=datetime.now(tz=UTC),
    )
    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=10,
        preview=QueryPreview(columns=["id"], rows=[[1]]),
    )
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(
            get=lambda execution_id: history,
            get_snapshot=lambda execution_id: snapshot,
        ),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=_CaptureArtifactService(),  # type: ignore[arg-type]
        snapshot_store=SimpleNamespace(iter_rows=lambda item: iter([[1], [2], [3]])),
        max_export_rows=2,
    )

    with pytest.raises(SqlResultExportLimitError) as error:
        service.export(
            execution_id="execution-large",
            export_format="xlsx",
            workspace_id="default",
            conversation_id="conversation-1",
            owner_user_id="user-1",
            allowed_current_execution_ids={"execution-large"},
        )

    assert error.value.code == "sql_export_row_limit"


def test_sql_export_rejects_output_before_exceeding_file_budget() -> None:
    history = SimpleNamespace(
        workspace_id="default",
        actor_id="user-1",
        status="succeeded",
        row_limit=10,
        preview=QueryPreview(columns=["value"], rows=[["x" * 200]]),
    )
    service = SqlResultExportService(
        execution_reader=SimpleNamespace(get=lambda execution_id: history),
        message_reader=SimpleNamespace(get_execution_message=lambda **kwargs: None),
        artifact_service=_CaptureArtifactService(),  # type: ignore[arg-type]
        max_export_bytes=64,
    )

    with pytest.raises(SqlResultExportLimitError) as error:
        service.export(
            execution_id="execution-small-budget",
            export_format="csv",
            workspace_id="default",
            conversation_id="conversation-1",
            owner_user_id="user-1",
            allowed_current_execution_ids={"execution-small-budget"},
        )

    assert error.value.code == "sql_export_size_limit"


def test_file_tools_are_web_only_and_have_explicit_schemas() -> None:
    names = ["list_files", "inspect_file", "run_file_code", "publish_files", "export_sql_result"]
    definitions = PromptBuilder.build_tools_from_names(names)
    assert [item["function"]["name"] for item in definitions] == names

    service = ToolFileCapabilityService(
        artifact_service=SimpleNamespace(),  # type: ignore[arg-type]
        sandbox_gateway=SimpleNamespace(),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )
    result = service.execute(
        tool_name="list_files",
        request=ChatAskRequest(
            question="test",
            actor_context=ChatActorContext(source="mattermost"),
        ),
        arguments={},
        state=ToolOrchestrationState(conversation_id="conversation-1"),
    )
    assert result["error_code"] == "file_tools_web_only"
    assert result["error_message"]
    assert result["can_continue"] is True
    assert result["retry_recommended"] is False


def test_sql_export_is_already_published_and_does_not_touch_sandbox() -> None:
    exported = FileArtifactRef(
        id="sql-artifact-1",
        filename="result.xlsx",
        media_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        size_bytes=128,
        kind="sql_result",
        sha256="a" * 64,
    )
    state = ToolOrchestrationState(conversation_id="conversation-1")
    state.published_files.append(exported)
    service = ToolFileCapabilityService(
        artifact_service=SimpleNamespace(),  # type: ignore[arg-type]
        # No sandbox methods exist: an accidental call fails this test immediately.
        sandbox_gateway=SimpleNamespace(),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )

    result = service.execute(
        tool_name="publish_files",
        request=ChatAskRequest(
            question="download",
            conversation_id="conversation-1",
            actor_context=ChatActorContext(source="web"),
        ),
        arguments={"output_ids": [exported.id]},
        state=state,
    )

    assert result == {
        "ok": True,
        "files": [exported.model_dump(mode="json")],
        "count": 0,
        "already_published": True,
        "requires_publish_files": False,
    }


def test_publish_files_is_idempotent_for_sandbox_output_id() -> None:
    artifact = Artifact(
        id="artifact-published",
        workspace_id="default",
        owner_user_id="user-1",
        conversation_id="conversation-1",
        storage_key="hidden/artifact-published",
        filename="result.csv",
        media_type="text/csv",
        size_bytes=5,
        sha256="a" * 64,
        source=ArtifactSource.SANDBOX_OUTPUT,
        status=ArtifactStatus.STAGED,
    )
    download_calls = 0

    def download_output(*, destination, **_kwargs):
        nonlocal download_calls
        download_calls += 1
        destination.write(b"id\n1\n")
        destination.seek(0)
        return "result.csv", "text/csv", 5

    service = ToolFileCapabilityService(
        artifact_service=SimpleNamespace(
            max_session_output_bytes=1024,
            max_output_bytes=1024,
            store_generated=lambda **_kwargs: artifact,
            to_ref=ArtifactService.to_ref,
        ),  # type: ignore[arg-type]
        sandbox_gateway=SimpleNamespace(download_output=download_output),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )
    request = ChatAskRequest(
        question="publish",
        workspace_id="default",
        conversation_id="conversation-1",
        actor_context=ChatActorContext(source="web"),
    )
    state = ToolOrchestrationState(
        conversation_id="conversation-1",
        sandbox_session_id="session-1",
        sandbox_outputs={
            "out-1": {
                "output_id": "out-1",
                "filename": "result.csv",
                "source_artifact_ids": [],
            }
        },
    )

    first = service.execute(
        tool_name="publish_files",
        request=request,
        arguments={"output_ids": ["out-1"]},
        state=state,
    )
    second = service.execute(
        tool_name="publish_files",
        request=request,
        arguments={"output_ids": ["out-1"]},
        state=state,
    )

    assert first["count"] == 1
    assert second == {
        "ok": True,
        "files": [ArtifactService.to_ref(artifact).model_dump(mode="json")],
        "count": 0,
        "already_published": True,
        "requires_publish_files": False,
    }
    assert download_calls == 1
    assert state.published_sandbox_outputs["out-1"].id == artifact.id


def test_file_tool_schema_explains_sql_export_delivery_contract() -> None:
    definitions = PromptBuilder.build_tools_from_names([
        "publish_files",
        "export_sql_result",
    ])
    descriptions = {
        item["function"]["name"]: item["function"]["description"] for item in definitions
    }

    assert "Never call this after export_sql_result" in descriptions["publish_files"]
    assert "do not call publish_files" in descriptions["export_sql_result"]


@pytest.mark.parametrize(
    ("tool_name", "tool_result", "required_keys"),
    [
        (
            "list_files",
            {
                "ok": True,
                "count": 1,
                "files": [
                    {
                        "id": "artifact-1",
                        "filename": "report.xlsx",
                        "media_type": (
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                        ),
                        "size_bytes": 128,
                        "kind": "user_upload",
                        "sha256": "a" * 64,
                    }
                ],
            },
            {"files", "count"},
        ),
        (
            "inspect_file",
            {
                "ok": True,
                "filename": "report.xlsx",
                "size_bytes": 128,
                "sheets": [{"name": "Sheet1", "rows": [["metric", "value"]]}],
            },
            {"filename", "size_bytes", "sheets"},
        ),
        (
            "run_file_code",
            {
                "ok": True,
                "exit_code": 0,
                "stdout": "created output",
                "stderr": "",
                "outputs": [
                    {
                        "output_id": "out-1",
                        "filename": "result.csv",
                        "media_type": "text/csv",
                        "size_bytes": 42,
                        "sha256": "b" * 64,
                    }
                ],
                "timed_out": False,
                "oom_killed": False,
            },
            {"exit_code", "stdout", "outputs", "timed_out", "oom_killed"},
        ),
        (
            "publish_files",
            {
                "ok": True,
                "count": 1,
                "files": [
                    {
                        "id": "artifact-2",
                        "filename": "result.csv",
                        "media_type": "text/csv",
                        "size_bytes": 42,
                        "kind": "sandbox_output",
                        "sha256": "b" * 64,
                    }
                ],
            },
            {"files", "count"},
        ),
    ],
)
def test_file_tool_results_keep_their_bounded_contract_in_the_llm_prompt(
    tool_name: str,
    tool_result: dict[str, object],
    required_keys: set[str],
) -> None:
    envelope = ToolResultPromptSerializer().serialize_envelope(
        tool_name=tool_name,
        tool_result=tool_result,
    )

    assert required_keys <= envelope.compact_payload.keys()
    assert envelope.compact_payload["ok"] is True


def test_inspect_file_preview_remains_bounded_but_usable_by_the_llm() -> None:
    preview = "x" * 12_000

    envelope = ToolResultPromptSerializer().serialize_envelope(
        tool_name="inspect_file",
        tool_result={"ok": True, "filename": "notes.txt", "preview": preview},
    )

    assert envelope.compact_payload["preview"] == preview


def test_failed_sandbox_process_returns_actionable_error_and_keeps_stderr() -> None:
    service = ToolFileCapabilityService(
        artifact_service=SimpleNamespace(),  # type: ignore[arg-type]
        sandbox_gateway=SimpleNamespace(
            run_python=lambda **_kwargs: SandboxRunResult(
                session_id="session-1",
                exit_code=1,
                stdout="",
                stderr="ModuleNotFoundError: missing_package",
            )
        ),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )

    result = service.execute(
        tool_name="run_file_code",
        request=ChatAskRequest(
            question="inspect",
            conversation_id="conversation-1",
            actor_context=ChatActorContext(source="web"),
        ),
        arguments={"code": "import missing_package", "artifact_ids": []},
        state=ToolOrchestrationState(conversation_id="conversation-1"),
    )

    assert result["error_code"] == "sandbox_process_failed"
    assert result["error_message"]
    assert result["stderr"] == "ModuleNotFoundError: missing_package"
    assert "Unknown error" not in StreamingResponseContentService.summarize_tool_result(result)


def test_http_sandbox_gateway_preserves_rejection_code_for_the_llm() -> None:
    def handle(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/internal/v1/files/sessions":
            return httpx.Response(200, json={"session_id": "session-1"})
        return httpx.Response(422, json={"detail": "inspection_failed"})

    gateway = HttpSandboxGateway(
        base_url="http://sandbox.test",
        internal_token="test-token",
        transport=httpx.MockTransport(handle),
    )

    with pytest.raises(SandboxGatewayError) as failure:
        gateway.inspect_file(
            session_id=None,
            input_file=SandboxInput(
                artifact_id="artifact-1",
                filename="report.xlsx",
                download_url="http://minio.test/report.xlsx",
                sha256="a" * 64,
            ),
        )

    assert failure.value.code == "sandbox_inspection_failed"
    assert failure.value.retry_recommended is False


def test_http_sandbox_gateway_cancels_session_after_timeout() -> None:
    requests: list[tuple[str, str]] = []

    def handle(request: httpx.Request) -> httpx.Response:
        requests.append((request.method, request.url.path))
        if request.url.path == "/internal/v1/files/sessions":
            return httpx.Response(200, json={"session_id": "session-timeout"})
        if request.url.path.endswith("/execute"):
            raise httpx.ReadTimeout("timed out", request=request)
        if request.method == "DELETE":
            return httpx.Response(204)
        raise AssertionError(request.url.path)

    gateway = HttpSandboxGateway(
        base_url="http://sandbox.test",
        internal_token="test-token",
        transport=httpx.MockTransport(handle),
    )

    with pytest.raises(SandboxGatewayError) as failure:
        gateway.run_python(session_id=None, code="print(1)", inputs=[])

    assert failure.value.code == "sandbox_timeout"
    assert ("DELETE", "/internal/v1/files/sessions/session-timeout") in requests


def test_file_tool_exposes_typed_sandbox_error_without_stopping_the_llm() -> None:
    def inspect_file(**_kwargs: object) -> dict[str, object]:
        raise SandboxGatewayError(
            "sandbox_input_hash_mismatch",
            "Sandbox rejected the operation: input_hash_mismatch.",
        )

    service = ToolFileCapabilityService(
        artifact_service=SimpleNamespace(
            get_owned=lambda **_kwargs: SimpleNamespace(conversation_id="conversation-1"),
            sandbox_input=lambda artifact: artifact,
        ),  # type: ignore[arg-type]
        sandbox_gateway=SimpleNamespace(inspect_file=inspect_file),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )

    result = service.execute(
        tool_name="inspect_file",
        request=ChatAskRequest(
            question="inspect",
            conversation_id="conversation-1",
            actor_context=ChatActorContext(source="web"),
        ),
        arguments={"artifact_id": "artifact-1"},
        state=ToolOrchestrationState(conversation_id="conversation-1"),
    )

    assert result == {
        "ok": False,
        "error_code": "sandbox_input_hash_mismatch",
        "error_message": "Sandbox rejected the operation: input_hash_mismatch.",
        "retry_recommended": False,
        "can_continue": True,
    }


def _file_session_runtime(tmp_path: Path) -> FileSessionRuntime:
    return FileSessionRuntime(
        FileSessionConfig(
            sessions_root=tmp_path / "sessions",
            workspace_mount=tmp_path / "workspace",
        ),
        ProcessExecutor(ProcessLimits(max_log_bytes=65536, pid_limit=64)),
    )


def test_execution_runtime_accepts_five_files_and_workflow_manifest(
    tmp_path: Path, monkeypatch
) -> None:
    supervisor = _file_session_runtime(tmp_path)
    session = supervisor._session(supervisor.create_session())

    def download(*, target, **_kwargs):
        target.write_bytes(b"{}")

    monkeypatch.setattr(supervisor, "_download", download)
    specs = [{"artifact_id": str(index), "filename": f"file-{index}.csv"} for index in range(5)] + [
        {"artifact_id": "manifest", "filename": "rhetorica-inputs.json"}
    ]
    paths = supervisor._materialize_inputs(session, specs)
    assert len(paths) == 6
    assert all(path.read_bytes() == b"{}" for path in paths)
    with pytest.raises(ExecutionRejectedError, match="too_many_inputs"):
        supervisor._materialize_inputs(session, [*specs, {"artifact_id": "extra"}])


def test_execution_output_collection_is_atomic_on_invalid_entry(tmp_path: Path) -> None:
    supervisor = _file_session_runtime(tmp_path)
    session = supervisor._session(supervisor.create_session())
    (session.root / "output" / "valid.csv").write_text("id\n1\n", encoding="utf-8")
    invalid = session.root / "output" / "invalid.exe"
    invalid.write_bytes(b"invalid")

    with pytest.raises(ValueError, match="unsupported_output_format"):
        supervisor._collect_outputs(session)

    assert session.outputs == {}
    assert session.output_versions == {}
    assert session.snapshot_bytes == 0
    assert not (session.root / "snapshots").exists()

    invalid.unlink()
    manifests = supervisor._collect_outputs(session)
    assert [item["filename"] for item in manifests] == ["valid.csv"]
    assert len(session.outputs) == 1


def test_sandbox_output_lineage_is_bound_to_the_run_that_created_it() -> None:
    artifacts = {
        artifact_id: Artifact(
            id=artifact_id,
            workspace_id="default",
            owner_user_id="user-1",
            conversation_id="conversation-1",
            storage_key=f"hidden/{artifact_id}",
            filename=f"{artifact_id}.csv",
            media_type="text/csv",
            size_bytes=1,
            sha256="a" * 64,
            source=ArtifactSource.USER_UPLOAD,
            status=ArtifactStatus.READY,
        )
        for artifact_id in ("input-a", "input-b")
    }
    artifact_service = SimpleNamespace(
        get_owned=lambda *, artifact_id, owner_user_id, workspace_id: artifacts[artifact_id],
        sandbox_input=lambda artifact: artifact,
    )
    run_number = 0

    def run_python(**_kwargs) -> SandboxRunResult:
        nonlocal run_number
        run_number += 1
        output_id = f"out-{run_number}"
        return SandboxRunResult(
            session_id="session-1",
            exit_code=0,
            stdout="",
            stderr="",
            outputs=(
                SandboxOutput(
                    output_id=output_id,
                    filename=f"result-{run_number}.csv",
                    media_type="text/csv",
                    size_bytes=1,
                    sha256=str(run_number) * 64,
                ),
            ),
        )

    service = ToolFileCapabilityService(
        artifact_service=artifact_service,  # type: ignore[arg-type]
        sandbox_gateway=SimpleNamespace(run_python=run_python),  # type: ignore[arg-type]
        sql_export_service=SimpleNamespace(),  # type: ignore[arg-type]
        owner_user_id="user-1",
    )
    request = ChatAskRequest(
        question="edit",
        workspace_id="default",
        conversation_id="conversation-1",
        actor_context=ChatActorContext(source="web"),
    )
    state = ToolOrchestrationState(conversation_id="conversation-1")

    first = service.execute(
        tool_name="run_file_code",
        request=request,
        arguments={"artifact_ids": ["input-a"], "code": "print('a')"},
        state=state,
    )
    second = service.execute(
        tool_name="run_file_code",
        request=request,
        arguments={"artifact_ids": ["input-b"], "code": "print('b')"},
        state=state,
    )

    assert first["outputs"][0]["output_id"] == "out-1"
    assert second["outputs"][0]["output_id"] == "out-2"
    assert "source_artifact_ids" not in second["outputs"][0]
    assert state.sandbox_outputs["out-1"]["source_artifact_ids"] == ["input-a"]
    assert state.sandbox_outputs["out-2"]["source_artifact_ids"] == ["input-b"]

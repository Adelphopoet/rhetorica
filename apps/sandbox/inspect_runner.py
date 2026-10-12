"""Bounded file inspector executed inside NsJail, never in the supervisor."""

from __future__ import annotations

import csv
import json
import sys
from itertools import islice
from pathlib import Path

_PREVIEW_CHARS = 12_000
_MAX_JSON_PARSE_BYTES = 2 * 1024 * 1024


def inspect(path: Path) -> dict[str, object]:
    suffix = path.suffix.lower()
    base: dict[str, object] = {"filename": path.name, "size_bytes": path.stat().st_size}
    if suffix in {".txt", ".md"}:
        with path.open(encoding="utf-8") as source:
            base["preview"] = source.read(_PREVIEW_CHARS)
    elif suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8", newline="") as source:
            rows = list(islice(
                csv.reader(source, delimiter="\t" if suffix == ".tsv" else ","),
                21,
            ))
        bounded_rows = [[_bounded_cell(cell) for cell in row[:20]] for row in rows]
        base.update({
            "columns": bounded_rows[0] if bounded_rows else [],
            "rows": bounded_rows[1:21],
        })
    elif suffix == ".json":
        with path.open(encoding="utf-8") as source:
            preview = source.read(_PREVIEW_CHARS)
        base["preview"] = preview
        if path.stat().st_size <= _MAX_JSON_PARSE_BYTES:
            payload = json.loads(path.read_text(encoding="utf-8"))
            base["structure"] = _structure(payload)
        else:
            base["structure"] = {"type": "large_json", "parsed": False}
    elif suffix == ".parquet":
        import pyarrow.parquet as pq
        parquet = pq.ParquetFile(path)
        first_batch = next(parquet.iter_batches(batch_size=10), None)
        rows = first_batch.to_pylist() if first_batch is not None else []
        base.update({
            "schema": str(parquet.schema_arrow)[:8_000],
            "row_count": parquet.metadata.num_rows,
            "rows": [
                {
                    str(key): _bounded_cell(value)
                    for key, value in list(row.items())[:20]
                }
                for row in rows
            ],
        })
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        base["sheets"] = [
            {
                "name": sheet.title,
                "rows": [
                    [_bounded_cell(cell) for cell in row[:15]]
                    for row in sheet.iter_rows(max_row=10, max_col=15, values_only=True)
                ],
            }
            for sheet in workbook.worksheets[:3]
        ]
        workbook.close()
    elif suffix == ".docx":
        from docx import Document
        document = Document(path)
        base["paragraphs"] = [_bounded_cell(p.text, 500) for p in document.paragraphs[:80]]
        base["table_count"] = len(document.tables)
    elif suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(path)
        base["slides"] = [
            [_bounded_cell(shape.text, 300) for shape in slide.shapes if hasattr(shape, "text")][
                :10
            ]
            for slide in presentation.slides[:10]
        ]
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        base["page_count"] = len(reader.pages)
        base["preview"] = "\n".join((page.extract_text() or "")[:4000] for page in reader.pages[:5])[:12000]
    elif suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image
        with Image.open(path) as image:
            base.update({"width": image.width, "height": image.height, "mode": image.mode, "format": image.format})
    return base


def _bounded_cell(value: object, limit: int = 120) -> object:
    if value is None or isinstance(value, int | float | bool):
        return value
    text = str(value)
    return text if len(text) <= limit else f"{text[:limit]}…"


def _structure(value: object, depth: int = 0) -> object:
    if depth >= 3:
        return type(value).__name__
    if isinstance(value, dict):
        return {str(key): _structure(item, depth + 1) for key, item in list(value.items())[:50]}
    if isinstance(value, list):
        return {"type": "array", "length": len(value), "sample": [_structure(item, depth + 1) for item in value[:3]]}
    return type(value).__name__


if __name__ == "__main__":
    print(json.dumps(inspect(Path(sys.argv[1])), ensure_ascii=False, default=str))
